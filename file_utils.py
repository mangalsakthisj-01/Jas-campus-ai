import asyncio
import mimetypes
import os
import subprocess
import tempfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from .ai import transcribe_media

ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".xml", ".html", ".css", ".js", ".ts",
    ".py", ".java", ".c", ".cpp", ".h", ".sql", ".sh", ".yaml", ".yml", ".toml", ".log",
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp",
    ".mp3", ".wav", ".m4a", ".mpga", ".ogg", ".flac", ".webm", ".mp4", ".mov", ".mkv", ".avi",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".m4a", ".mpga", ".ogg", ".flac"}
VIDEO_EXTENSIONS = {".webm", ".mp4", ".mov", ".mkv", ".avi"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".html", ".css", ".js", ".ts", ".py", ".java", ".c", ".cpp",
    ".h", ".sql", ".sh", ".yaml", ".yml", ".toml", ".log"
}

def infer_mime(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXTENSIONS:
        return mimetypes.guess_type(filename)[0] or "application/octet-stream"
    if ext in AUDIO_EXTENSIONS:
        return mimetypes.guess_type(filename)[0] or "audio/mpeg"
    if ext in VIDEO_EXTENSIONS:
        return mimetypes.guess_type(filename)[0] or "video/mp4"
    return mimetypes.guess_type(filename)[0] or "application/octet-stream"

def extract_text(path: Path, extension: str) -> str:
    ext = extension.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(str(path))
            sections = []
            for i, page in enumerate(reader.pages):
                content = page.extract_text() or ""
                if content.strip():
                    sections.append(f"[Page {i + 1}]\n{content}")
            text = "\n\n".join(sections)
            return text[:80000] if text else "[No selectable text found in this PDF. It may be scanned; OCR is not enabled in this starter. Try a text-based PDF or enable OCR separately.]"
        if ext == ".docx":
            doc = Document(str(path))
            chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(chunks)[:80000]
        if ext in TEXT_EXTENSIONS:
            return path.read_text(encoding="utf-8", errors="replace")[:80000]
        if ext in IMAGE_EXTENSIONS:
            return "[Image attachment. The image itself is sent to the model when a vision-capable model is configured.]"
        if ext in AUDIO_EXTENSIONS or ext in VIDEO_EXTENSIONS:
            return "[Media received. Transcription is being attempted if the OpenAI transcription API is configured.]"
        return "[File stored; text extraction is not supported for this file type.]"
    except Exception as exc:
        return f"[Could not extract text from this file: {type(exc).__name__}. The original upload is stored.]"

async def process_media(path: Path, extension: str) -> str:
    ext = extension.lower()
    try:
        if ext in VIDEO_EXTENSIONS:
            # Fixed command arguments (no shell) prevent prompt or filename contents becoming shell instructions.
            with tempfile.NamedTemporaryFile(prefix="jas-audio-", suffix=".mp3", delete=False) as tmp:
                audio_path = Path(tmp.name)
            try:
                proc = await asyncio.create_subprocess_exec(
                    "ffmpeg", "-nostdin", "-y", "-i", str(path), "-vn", "-ac", "1", "-ar", "16000", "-b:a", "64k", str(audio_path),
                    stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL,
                )
                try:
                    await asyncio.wait_for(proc.communicate(), timeout=180)
                except asyncio.TimeoutError:
                    proc.kill()
                    await proc.communicate()
                    return "[Video saved, but audio extraction timed out. Install ffmpeg and try a shorter file.]"
                if proc.returncode != 0 or not audio_path.exists() or audio_path.stat().st_size == 0:
                    return "[Video saved, but audio could not be extracted. Ensure ffmpeg is installed and the video contains an audio track.]"
                if audio_path.stat().st_size > 24 * 1024 * 1024:
                    return "[Video saved, but extracted audio is too large for the configured transcription path. Trim the clip and upload again.]"
                return await transcribe_media(audio_path) or "[No speech was detected in the video.]"
            finally:
                audio_path.unlink(missing_ok=True)
        if ext in AUDIO_EXTENSIONS:
            if path.stat().st_size > 24 * 1024 * 1024:
                return "[Audio saved, but it is too large for the configured transcription path. Trim the clip and upload again.]"
            return await transcribe_media(path) or "[No speech was detected in the audio.]"
    except FileNotFoundError:
        return "[Media saved, but ffmpeg is not installed. Install ffmpeg to process video files.]"
    except RuntimeError as exc:
        return f"[{exc}]"
    except Exception as exc:
        return f"[Media saved, but transcription failed ({type(exc).__name__}).]"
    return ""
