import base64
import json
import os
import re
import secrets
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated

from dotenv import load_dotenv

# Load project-local .env for the simple non-Docker setup; existing environment wins.
load_dotenv()

from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.staticfiles import StaticFiles
from sqlalchemy import desc, func, select
from sqlalchemy.orm import Session

from .ai import generate_chat, generate_image
from .database import Base, SessionLocal, engine, get_db
from .file_utils import ALLOWED_EXTENSIONS, IMAGE_EXTENSIONS, extract_text, infer_mime, process_media
from .models import AuditLog, ChatMessage, Conversation, Upload, User, utcnow
from .schemas import ChatRequest, DocumentRequest, ImageRequest, LoginRequest, ManagerCreateRequest, RegisterRequest, UserDecisionRequest
from .security import create_access_token, decode_access_token, hash_password, verify_password

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
STORAGE_DIR = Path(os.getenv("JAS_STORAGE_DIR", "./storage")).resolve()
UPLOAD_DIR = STORAGE_DIR / "uploads"
GENERATED_DIR = STORAGE_DIR / "generated"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
GENERATED_DIR.mkdir(parents=True, exist_ok=True)
MAX_UPLOAD_MB = int(os.getenv("JAS_MAX_UPLOAD_MB", "25"))
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024
ADMIN_CAN_VIEW_UPLOAD_CONTENT = os.getenv("JAS_ADMIN_CAN_VIEW_UPLOAD_CONTENT", "true").lower() in {"1", "true", "yes"}
ALLOWED_EMAIL_DOMAINS = {
    domain.strip().lower().lstrip("@")
    for domain in os.getenv("JAS_ALLOWED_EMAIL_DOMAINS", "").split(",")
    if domain.strip()
}
CORS_ORIGINS = [origin.strip() for origin in os.getenv("JAS_CORS_ORIGINS", "").split(",") if origin.strip()]

app = FastAPI(
    title="JAS Campus AI",
    version="0.1.0",
    description="A role-controlled campus AI assistant starter application.",
)
if CORS_ORIGINS:
    app.add_middleware(
        CORSMiddleware,
        allow_origins=CORS_ORIGINS,
        allow_credentials=False,
        allow_methods=["GET", "POST", "DELETE", "OPTIONS"],
        allow_headers=["Authorization", "Content-Type"],
    )
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
bearer_scheme = HTTPBearer(auto_error=False)


def audit(db: Session, actor_id: int | None, action: str, target_type: str = "", target_id: str | int = "", details: dict | None = None):
    db.add(AuditLog(
        actor_id=actor_id,
        action=action,
        target_type=target_type,
        target_id=str(target_id),
        details=json.dumps(details or {}, ensure_ascii=False),
    ))


def enforce_allowed_email_domain(email: str) -> None:
    if not ALLOWED_EMAIL_DOMAINS:
        return
    domain = email.rsplit("@", 1)[-1].lower()
    if domain not in ALLOWED_EMAIL_DOMAINS:
        allowed = ", ".join(sorted(ALLOWED_EMAIL_DOMAINS))
        raise HTTPException(status_code=403, detail=f"JAS registration is restricted to approved college email domains: {allowed}.")


def serialize_user(user: User) -> dict:
    return {
        "id": user.id,
        "email": user.email,
        "full_name": user.full_name,
        "role": user.role,
        "status": user.status,
        "created_at": user.created_at.isoformat() if user.created_at else None,
    }


def serialize_upload(upload: Upload, include_text: bool = False) -> dict:
    result = {
        "id": upload.id,
        "owner_id": upload.owner_id,
        "filename": upload.filename,
        "mime_type": upload.mime_type,
        "extension": upload.extension,
        "size_bytes": upload.size_bytes,
        "processing_status": upload.processing_status,
        "created_at": upload.created_at.isoformat() if upload.created_at else None,
    }
    if include_text:
        result["extracted_text"] = upload.extracted_text
    return result


def current_user(
    credentials: Annotated[HTTPAuthorizationCredentials | None, Depends(bearer_scheme)],
    db: Annotated[Session, Depends(get_db)],
) -> User:
    if not credentials:
        raise HTTPException(status_code=401, detail="Sign in to continue.", headers={"WWW-Authenticate": "Bearer"})
    user_id = decode_access_token(credentials.credentials)
    user = db.get(User, user_id) if user_id is not None else None
    if user is None:
        raise HTTPException(status_code=401, detail="Your session is invalid or expired. Sign in again.")
    return user


def approved_user(user: Annotated[User, Depends(current_user)]) -> User:
    if user.status != "approved":
        raise HTTPException(status_code=403, detail=f"Account status is {user.status}. A college admin or manager must approve access first.")
    return user


def role_required(*roles: str):
    def dependency(user: Annotated[User, Depends(approved_user)]) -> User:
        if user.role not in roles:
            raise HTTPException(status_code=403, detail="You do not have permission to perform this action.")
        return user
    return dependency


def get_owned_conversation(db: Session, conversation_id: int, user_id: int) -> Conversation:
    conversation = db.get(Conversation, conversation_id)
    if not conversation or conversation.owner_id != user_id:
        raise HTTPException(status_code=404, detail="Conversation not found.")
    return conversation


def attachment_context(upload: Upload) -> str:
    if upload.extension in IMAGE_EXTENSIONS:
        return f"\n\n[Image attached: {upload.filename}. The image was supplied separately to the model for visual analysis.]"
    text = (upload.extracted_text or "").strip()
    if not text:
        return f"\n\n[Attachment: {upload.filename}. No extractable text is available.]"
    return f"\n\n[Attachment: {upload.filename}]\n{text[:18000]}\n[End attachment]"


def build_model_user_content(message: str, attachments: list[Upload]) -> str | list[dict]:
    text = message + "\n\n" + "\n".join(attachment_context(item) for item in attachments) if attachments else message
    images = []
    total_image_bytes = 0
    for item in attachments:
        if item.extension not in IMAGE_EXTENSIONS:
            continue
        path = UPLOAD_DIR / item.stored_name
        if not path.exists():
            continue
        if item.size_bytes > 10 * 1024 * 1024:
            text += f"\n\n[Image {item.filename} was not sent to the model because it exceeds the 10 MB per-image limit.]"
            continue
        total_image_bytes += item.size_bytes
        if total_image_bytes > 18 * 1024 * 1024:
            text += f"\n\n[Image {item.filename} was not sent because combined image size exceeds the limit.]"
            continue
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        mime = item.mime_type if item.mime_type.startswith("image/") else "image/png"
        images.append({"type": "image_url", "image_url": {"url": f"data:{mime};base64,{encoded}"}})
    if not images:
        return text
    return [{"type": "text", "text": text}] + images


@app.on_event("startup")
def startup():
    production = os.getenv("JAS_ENV", "development").lower() == "production"
    if production:
        secret = os.getenv("JAS_JWT_SECRET", "")
        if len(secret) < 32 or "replace" in secret.lower() or "change-this" in secret.lower():
            raise RuntimeError("Production startup blocked: set JAS_JWT_SECRET to a unique random secret of at least 32 characters.")
        db_password = os.getenv("POSTGRES_PASSWORD", "")
        if not db_password or "replace" in db_password.lower():
            raise RuntimeError("Production startup blocked: set POSTGRES_PASSWORD to a unique strong password.")
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        existing_admin = db.scalar(select(User).where(User.role == "admin", User.status == "approved"))
        if existing_admin:
            return
        email = os.getenv("JAS_BOOTSTRAP_ADMIN_EMAIL", "").strip().lower()
        password = os.getenv("JAS_BOOTSTRAP_ADMIN_PASSWORD", "")
        if not email or not password:
            if production:
                raise RuntimeError("No approved admin exists. Set JAS_BOOTSTRAP_ADMIN_EMAIL and JAS_BOOTSTRAP_ADMIN_PASSWORD for initial setup.")
            return
        if production and (len(password) < 12 or "change-this" in password.lower()):
            raise RuntimeError("Production startup blocked: set JAS_BOOTSTRAP_ADMIN_PASSWORD to a unique password (12+ characters).")
        existing = db.scalar(select(User).where(User.email == email))
        if existing:
            if existing.role != "admin" or existing.status != "approved":
                raise RuntimeError("Bootstrap email already exists but is not an approved admin. Choose an unused bootstrap email or repair the account manually.")
            return
        name = os.getenv("JAS_BOOTSTRAP_ADMIN_NAME", "JAS Administrator").strip() or "JAS Administrator"
        admin = User(email=email, full_name=name, password_hash=hash_password(password), role="admin", status="approved")
        db.add(admin)
        db.flush()
        audit(db, admin.id, "bootstrap_admin_created", "user", admin.id)
        db.commit()
    finally:
        db.close()

@app.get("/", response_class=HTMLResponse)
def home():
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "JAS Campus AI", "version": "0.1.0"}


@app.post("/api/auth/register", status_code=201)
def register(payload: RegisterRequest, db: Annotated[Session, Depends(get_db)]):
    email = str(payload.email).lower().strip()
    enforce_allowed_email_domain(email)
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(status_code=409, detail="An account with this email already exists.")
    user = User(
        email=email,
        full_name=payload.full_name,
        password_hash=hash_password(payload.password),
        role="user",
        status="pending",
    )
    db.add(user)
    db.flush()
    audit(db, user.id, "user_registered", "user", user.id, {"role": "user"})
    db.commit()
    return {"message": "Registration received. A college admin or manager must approve your account before you can use JAS."}


@app.post("/api/auth/login")
def login(payload: LoginRequest, db: Annotated[Session, Depends(get_db)]):
    email = str(payload.email).lower().strip()
    user = db.scalar(select(User).where(User.email == email))
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    if user.status == "disabled":
        raise HTTPException(status_code=403, detail="This account is disabled. Contact the college JAS administrator.")
    audit(db, user.id, "user_login", "user", user.id)
    db.commit()
    return {"access_token": create_access_token(user.id), "token_type": "bearer", "user": serialize_user(user)}


@app.get("/api/me")
def me(user: Annotated[User, Depends(current_user)]):
    return {"user": serialize_user(user), "max_upload_mb": MAX_UPLOAD_MB}


@app.get("/api/conversations")
def list_conversations(user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    rows = db.scalars(select(Conversation).where(Conversation.owner_id == user.id).order_by(desc(Conversation.updated_at)).limit(100)).all()
    return [{"id": c.id, "title": c.title, "updated_at": c.updated_at.isoformat()} for c in rows]


@app.post("/api/conversations", status_code=201)
def create_conversation(user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    conversation = Conversation(owner_id=user.id, title="New chat")
    db.add(conversation)
    db.flush()
    audit(db, user.id, "conversation_created", "conversation", conversation.id)
    db.commit()
    return {"id": conversation.id, "title": conversation.title, "updated_at": conversation.updated_at.isoformat()}


@app.get("/api/conversations/{conversation_id}/messages")
def list_messages(conversation_id: int, user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    conversation = get_owned_conversation(db, conversation_id, user.id)
    rows = db.scalars(select(ChatMessage).where(ChatMessage.conversation_id == conversation.id).order_by(ChatMessage.created_at.asc()).limit(500)).all()
    return [{"id": m.id, "role": m.role, "content": m.content, "created_at": m.created_at.isoformat()} for m in rows]


@app.post("/api/chat")
async def chat(payload: ChatRequest, user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    if payload.conversation_id:
        conversation = get_owned_conversation(db, payload.conversation_id, user.id)
    else:
        conversation = Conversation(owner_id=user.id, title=" ".join(payload.message.split())[:72] or "New chat")
        db.add(conversation)
        db.flush()

    attachment_ids = list(dict.fromkeys(payload.attachment_ids))
    attachments: list[Upload] = []
    if attachment_ids:
        rows = db.scalars(select(Upload).where(Upload.id.in_(attachment_ids), Upload.owner_id == user.id)).all()
        by_id = {item.id: item for item in rows}
        missing = [item for item in attachment_ids if item not in by_id]
        if missing:
            raise HTTPException(status_code=404, detail="One or more attachments were not found in your uploads.")
        attachments = [by_id[item] for item in attachment_ids]

    previous = db.scalars(
        select(ChatMessage).where(ChatMessage.conversation_id == conversation.id).order_by(desc(ChatMessage.id)).limit(16)
    ).all()
    history = []
    for item in reversed(previous):
        content = item.content
        try:
            past_ids = json.loads(item.attachment_ids or "[]")
        except json.JSONDecodeError:
            past_ids = []
        if item.role == "user" and past_ids:
            past_uploads = db.scalars(select(Upload).where(Upload.id.in_(past_ids), Upload.owner_id == user.id)).all()
            content += "\n" + "\n".join(attachment_context(u) for u in past_uploads)
        history.append({"role": item.role, "content": content})

    user_content = build_model_user_content(payload.message, attachments)
    history.append({"role": "user", "content": user_content})

    user_message = ChatMessage(conversation_id=conversation.id, role="user", content=payload.message, attachment_ids=json.dumps(attachment_ids))
    db.add(user_message)
    if conversation.title == "New chat" and payload.message.strip():
        conversation.title = " ".join(payload.message.split())[:72]
    conversation.updated_at = utcnow()
    audit(db, user.id, "chat_requested", "conversation", conversation.id, {"mode": payload.mode, "attachment_count": len(attachments), "message_length": len(payload.message)})
    db.commit()

    try:
        answer = await generate_chat(history, payload.mode)
    except Exception as exc:
        # Do not leak provider secrets or upstream request details to users.
        error_text = str(exc)
        if "No model API key configured" in error_text or "No API key configured" in error_text:
            raise HTTPException(status_code=503, detail="JAS has no model configured yet. Set JAS_LLM_API_KEY and JAS_LLM_MODEL, or connect a local OpenAI-compatible model server.")
        raise HTTPException(status_code=502, detail="The AI provider could not complete this request. Check the model configuration and try again.")

    assistant_message = ChatMessage(conversation_id=conversation.id, role="assistant", content=answer, attachment_ids="[]")
    db.add(assistant_message)
    conversation.updated_at = utcnow()
    audit(db, user.id, "chat_response_created", "conversation", conversation.id, {"answer_length": len(answer), "mode": payload.mode})
    db.commit()
    return {
        "conversation_id": conversation.id,
        "title": conversation.title,
        "message": {"id": assistant_message.id, "role": "assistant", "content": answer, "created_at": assistant_message.created_at.isoformat()},
    }


@app.post("/api/uploads", status_code=201)
async def upload_file(file: Annotated[UploadFile, File(...)], user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    safe_filename = Path(file.filename or "upload").name[:255]
    ext = Path(safe_filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail=f"File type not supported. Allowed extensions: {', '.join(sorted(ALLOWED_EXTENSIONS))}")
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"File exceeds the {MAX_UPLOAD_MB} MB upload limit.")
    if not raw:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    stored_name = secrets.token_hex(20) + ext
    path = UPLOAD_DIR / stored_name
    path.write_bytes(raw)
    mime = infer_mime(safe_filename)
    extracted = extract_text(path, ext)
    processing_status = "processed"
    if ext in {".mp3", ".wav", ".m4a", ".mpga", ".ogg", ".flac", ".webm", ".mp4", ".mov", ".mkv", ".avi"}:
        media_text = await process_media(path, ext)
        if media_text:
            extracted = media_text[:80000]
        processing_status = "transcribed" if extracted and not extracted.startswith("[") else "stored"
    upload = Upload(
        owner_id=user.id,
        filename=safe_filename,
        stored_name=stored_name,
        mime_type=mime,
        extension=ext,
        size_bytes=len(raw),
        extracted_text=extracted,
        processing_status=processing_status,
    )
    db.add(upload)
    db.flush()
    audit(db, user.id, "file_uploaded", "upload", upload.id, {"filename": safe_filename, "size_bytes": len(raw), "extension": ext})
    db.commit()
    return serialize_upload(upload)


@app.get("/api/uploads")
def my_uploads(user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    rows = db.scalars(select(Upload).where(Upload.owner_id == user.id).order_by(desc(Upload.created_at)).limit(100)).all()
    return [serialize_upload(item) for item in rows]


@app.delete("/api/uploads/{upload_id}")
def delete_upload(upload_id: int, user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    upload = db.get(Upload, upload_id)
    if not upload or upload.owner_id != user.id:
        raise HTTPException(status_code=404, detail="Upload not found.")
    (UPLOAD_DIR / upload.stored_name).unlink(missing_ok=True)
    audit(db, user.id, "upload_deleted", "upload", upload.id, {"filename": upload.filename})
    db.delete(upload)
    db.commit()
    return {"message": "Upload deleted."}


@app.get("/api/uploads/{upload_id}/download")
def download_my_upload(upload_id: int, user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    upload = db.get(Upload, upload_id)
    if not upload or upload.owner_id != user.id:
        raise HTTPException(status_code=404, detail="Upload not found.")
    path = UPLOAD_DIR / upload.stored_name
    if not path.exists():
        raise HTTPException(status_code=404, detail="Stored file is unavailable.")
    return FileResponse(path, filename=upload.filename, media_type=upload.mime_type)


@app.post("/api/exports/document")
def export_document(payload: DocumentRequest, user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    safe_title = re.sub(r"[^A-Za-z0-9_-]+", "_", payload.title).strip("_")[:64] or "JAS_Document"
    filename = f"{safe_title}_{secrets.token_hex(6)}.{payload.format}"
    path = GENERATED_DIR / filename
    if payload.format == "docx":
        from docx import Document
        from docx.shared import Pt
        document = Document()
        document.add_heading(payload.title, level=0)
        in_code = False
        for line in payload.content.splitlines():
            stripped = line.strip()
            if stripped.startswith("```"):
                in_code = not in_code
                continue
            if not stripped:
                document.add_paragraph("")
            elif not in_code and stripped.startswith("#"):
                heading = stripped.lstrip("# ").strip()
                level = min(len(stripped) - len(stripped.lstrip("#")), 3)
                document.add_heading(heading or "Section", level=max(1, level))
            elif not in_code and stripped.startswith(("- ", "* ")):
                document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            elif not in_code and re.match(r"^\d+[.)]\s+", stripped):
                document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number")
            else:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(line if in_code else stripped)
                if in_code:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
        document.save(path)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        from xml.sax.saxutils import escape
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        styles = getSampleStyleSheet()
        code_style = ParagraphStyle("JASCode", parent=styles["Code"], fontName="Courier", fontSize=8, leading=10, backColor=colors.whitesmoke, leftIndent=6, rightIndent=6, borderPadding=4)
        story = [Paragraph(escape(payload.title), styles["Title"]), Spacer(1, 5 * mm)]
        in_code = False
        for line in payload.content.splitlines() or [""]:
            stripped = line.strip()
            if stripped.startswith("```"):
                in_code = not in_code
                story.append(Spacer(1, 2 * mm))
                continue
            if not stripped:
                story.append(Spacer(1, 3 * mm))
            elif in_code:
                story.append(Paragraph(escape(line) or "&nbsp;", code_style))
            elif stripped.startswith("#"):
                heading = stripped.lstrip("# ").strip() or "Section"
                level = min(len(stripped) - len(stripped.lstrip("#")), 3)
                story.append(Paragraph(escape(heading), styles[f"Heading{max(1, level)}"]))
                story.append(Spacer(1, 1.5 * mm))
            elif stripped.startswith(("- ", "* ")):
                story.append(Paragraph(escape(stripped[2:].strip()), styles["BodyText"], bulletText="•"))
                story.append(Spacer(1, 1 * mm))
            elif re.match(r"^\d+[.)]\s+", stripped):
                story.append(Paragraph(escape(stripped), styles["BodyText"]))
                story.append(Spacer(1, 1 * mm))
            else:
                story.append(Paragraph(escape(line), styles["BodyText"]))
                story.append(Spacer(1, 1.5 * mm))
        pdf = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=payload.title, author="JAS Campus AI")
        pdf.build(story)
        media_type = "application/pdf"
    audit(db, user.id, "document_exported", "document", filename, {"format": payload.format, "title": payload.title, "content_length": len(payload.content)})
    db.commit()
    return FileResponse(path, filename=filename, media_type=media_type)


@app.post("/api/generate-image")
async def create_image(payload: ImageRequest, user: Annotated[User, Depends(approved_user)], db: Annotated[Session, Depends(get_db)]):
    try:
        data_url = await generate_image(payload.prompt, user.email)
    except Exception as exc:
        if "Image generation is not configured" in str(exc):
            raise HTTPException(status_code=503, detail="Image generation is not configured. Set OPENAI_API_KEY and JAS_IMAGE_MODEL.")
        raise HTTPException(status_code=502, detail="The image provider could not generate an image. Verify model access and configuration.")
    audit(db, user.id, "image_generated", "image", "", {"prompt_length": len(payload.prompt)})
    db.commit()
    return {"image_data_url": data_url}


@app.get("/api/manage/users")
def manage_users(user: Annotated[User, Depends(role_required("admin", "manager"))], db: Annotated[Session, Depends(get_db)]):
    rows = db.scalars(select(User).order_by(User.created_at.desc()).limit(500)).all()
    return [serialize_user(row) for row in rows]


@app.post("/api/manage/users/{user_id}/decision")
def decide_user(user_id: int, payload: UserDecisionRequest, user: Annotated[User, Depends(role_required("admin", "manager"))], db: Annotated[Session, Depends(get_db)]):
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(status_code=404, detail="User not found.")
    if target.id == user.id:
        raise HTTPException(status_code=400, detail="You cannot change your own account status.")
    if target.role == "admin":
        raise HTTPException(status_code=403, detail="Admin accounts cannot be changed from this screen.")
    if user.role == "manager" and target.role != "user":
        raise HTTPException(status_code=403, detail="Managers may manage student/user accounts only. Only an admin can manage managers.")
    target.status = payload.status
    target.approved_by_id = user.id if payload.status == "approved" else target.approved_by_id
    audit(db, user.id, f"user_status_{payload.status}", "user", target.id, {"target_email": target.email, "target_role": target.role})
    db.commit()
    return serialize_user(target)


@app.post("/api/admin/managers", status_code=201)
def create_manager(payload: ManagerCreateRequest, user: Annotated[User, Depends(role_required("admin"))], db: Annotated[Session, Depends(get_db)]):
    email = str(payload.email).lower().strip()
    enforce_allowed_email_domain(email)
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(status_code=409, detail="An account with this email already exists.")
    manager = User(email=email, full_name=payload.full_name, password_hash=hash_password(payload.password), role="manager", status="approved", approved_by_id=user.id)
    db.add(manager)
    db.flush()
    audit(db, user.id, "manager_created", "user", manager.id, {"manager_email": email})
    db.commit()
    return serialize_user(manager)


@app.get("/api/admin/overview")
def admin_overview(user: Annotated[User, Depends(role_required("admin"))], db: Annotated[Session, Depends(get_db)]):
    count = lambda model, *criteria: db.scalar(select(func.count()).select_from(model).where(*criteria)) or 0
    return {
        "users_total": count(User),
        "users_pending": count(User, User.status == "pending"),
        "users_approved": count(User, User.status == "approved"),
        "managers_total": count(User, User.role == "manager"),
        "uploads_total": count(Upload),
        "conversations_total": count(Conversation),
    }


@app.get("/api/admin/uploads")
def admin_uploads(user: Annotated[User, Depends(role_required("admin"))], db: Annotated[Session, Depends(get_db)], limit: int = 100):
    limit = max(1, min(limit, 200))
    rows = db.scalars(select(Upload).order_by(desc(Upload.created_at)).limit(limit)).all()
    result = []
    for item in rows:
        owner = db.get(User, item.owner_id)
        row = serialize_upload(item, include_text=ADMIN_CAN_VIEW_UPLOAD_CONTENT)
        row["owner_email"] = owner.email if owner else "unknown"
        row["owner_name"] = owner.full_name if owner else "Unknown"
        result.append(row)
    audit(db, user.id, "admin_upload_review_list", "uploads", "", {"count": len(rows)})
    db.commit()
    return {"content_review_enabled": ADMIN_CAN_VIEW_UPLOAD_CONTENT, "uploads": result}


@app.get("/api/admin/uploads/{upload_id}/download")
def admin_download_upload(upload_id: int, user: Annotated[User, Depends(role_required("admin"))], db: Annotated[Session, Depends(get_db)]):
    upload = db.get(Upload, upload_id)
    if not upload:
        raise HTTPException(status_code=404, detail="Upload not found.")
    path = UPLOAD_DIR / upload.stored_name
    if not path.exists():
        raise HTTPException(status_code=404, detail="Stored file is unavailable.")
    audit(db, user.id, "admin_upload_download", "upload", upload.id, {"owner_id": upload.owner_id})
    db.commit()
    return FileResponse(path, filename=upload.filename, media_type=upload.mime_type)


@app.get("/api/admin/audit-logs")
def admin_audit_logs(user: Annotated[User, Depends(role_required("admin"))], db: Annotated[Session, Depends(get_db)], limit: int = 100):
    limit = max(1, min(limit, 300))
    rows = db.scalars(select(AuditLog).order_by(desc(AuditLog.created_at)).limit(limit)).all()
    return [{
        "id": item.id,
        "actor_id": item.actor_id,
        "action": item.action,
        "target_type": item.target_type,
        "target_id": item.target_id,
        "details": item.details,
        "created_at": item.created_at.isoformat(),
    } for item in rows]
